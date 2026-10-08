from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document

from .pdf_parser import PDFParseError, ParsedPDF, parse_pdf


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
}


@dataclass
class LoadedResume:
    """
    Normalized representation of one successfully loaded resume.
    """

    filename: str
    path: str
    file_type: str
    text: str
    page_count: int | None = None


@dataclass
class ResumeLoadFailure:
    """
    Represents a resume that could not be loaded.

    This is deliberately separate from candidate rejection.

    FAILED:
        We couldn't reliably read the resume.

    REJECTED:
        We successfully read the resume but the candidate
        did not satisfy eligibility requirements.
    """

    filename: str
    path: str
    stage: str
    error_type: str
    message: str


def normalize_document_text(text: str) -> str:
    """
    Apply conservative text normalization.

    We preserve line breaks because they can contain useful
    structural information in resumes.
    """

    if not text:
        return ""

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", "")

    normalized_lines: list[str] = []
    previous_blank = False

    for raw_line in text.split("\n"):
        line = raw_line.strip()

        if not line:
            if not previous_blank:
                normalized_lines.append("")

            previous_blank = True
            continue

        normalized_lines.append(line)
        previous_blank = False

    return "\n".join(normalized_lines).strip()


def load_pdf(path: Path) -> LoadedResume:
    """
    Load a PDF resume using the dedicated PDF parser.
    """

    parsed: ParsedPDF = parse_pdf(path)

    return LoadedResume(
        filename=parsed.filename,
        path=parsed.path,
        file_type="pdf",
        text=parsed.text,
        page_count=parsed.page_count,
    )


def load_docx(path: Path) -> LoadedResume:
    """
    Load a DOCX resume.

    DOCX support is a bonus capability, but implementing it here
    costs very little and makes the ingestion layer more robust.
    """

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ValueError(
            f"Could not open DOCX: {exc}"
        ) from exc

    blocks: list[str] = []

    # Paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            blocks.append(text)

    # Tables are common in resumes.
    #
    # Important:
    # We don't assume that the resume uses paragraphs only.
    # Skills, education, contact information, etc. may be inside
    # tables.
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                blocks.append(" | ".join(cells))

    full_text = normalize_document_text(
        "\n".join(blocks)
    )

    if not full_text:
        raise ValueError(
            "DOCX contains no extractable text."
        )

    return LoadedResume(
        filename=path.name,
        path=str(path),
        file_type="docx",
        text=full_text,
        page_count=None,
    )


def load_txt(path: Path) -> LoadedResume:
    """
    Load a plain-text resume.

    UTF-8 is attempted first. If a resume was saved using another
    common encoding, we fall back without crashing the batch.
    """

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1",
    ]

    last_error: Exception | None = None

    for encoding in encodings:
        try:
            text = path.read_text(
                encoding=encoding
            )

            text = normalize_document_text(text)

            if not text:
                raise ValueError(
                    "TXT file contains no extractable text."
                )

            return LoadedResume(
                filename=path.name,
                path=str(path),
                file_type="txt",
                text=text,
                page_count=None,
            )

        except UnicodeDecodeError as exc:
            last_error = exc

    raise ValueError(
        "Could not decode TXT file using supported encodings."
    ) from last_error


def load_resume(path: Path) -> LoadedResume:
    """
    Load one resume according to its file extension.

    Unsupported formats are rejected at the ingestion stage.
    """

    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {suffix or '[no extension]'}"
        )

    if suffix == ".pdf":
        return load_pdf(path)

    if suffix == ".docx":
        return load_docx(path)

    if suffix == ".txt":
        return load_txt(path)

    # Defensive fallback.
    raise ValueError(
        f"No loader implemented for file type: {suffix}"
    )


def discover_resume_files(
    input_dir: Path,
) -> list[Path]:
    """
    Discover supported resume files recursively.

    Recursive discovery means that if the user later places resumes
    inside subfolders, the pipeline can still find them.
    """

    if not input_dir.exists():
        raise FileNotFoundError(
            f"Resume directory does not exist: {input_dir}"
        )

    if not input_dir.is_dir():
        raise NotADirectoryError(
            f"Resume path is not a directory: {input_dir}"
        )

    files = [
        path
        for path in input_dir.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    return sorted(
        files,
        key=lambda path: str(path).lower(),
    )


def load_all_resumes(
    input_dir: Path,
) -> tuple[list[LoadedResume], list[ResumeLoadFailure]]:
    """
    Load every supported resume in the directory.

    IMPORTANT:
    A failure in one resume does not stop processing of the others.
    """

    resume_files = discover_resume_files(input_dir)

    loaded: list[LoadedResume] = []
    failures: list[ResumeLoadFailure] = []

    for path in resume_files:
        try:
            resume = load_resume(path)

            loaded.append(resume)

        except PDFParseError as exc:
            failures.append(
                ResumeLoadFailure(
                    filename=path.name,
                    path=str(path),
                    stage="ingestion",
                    error_type="PDFParseError",
                    message=str(exc),
                )
            )

        except Exception as exc:
            failures.append(
                ResumeLoadFailure(
                    filename=path.name,
                    path=str(path),
                    stage="ingestion",
                    error_type=type(exc).__name__,
                    message=str(exc),
                )
            )

    return loaded, failures